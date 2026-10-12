from __future__ import annotations

import csv
import io
import math
import unicodedata
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


SUPPORTED_EXTENSIONS = {
    ".xlsx",
    ".xls",
    ".xlsm",
    ".xlsb",
    ".xltx",
    ".xltm",
    ".xlt",
    ".xml",
    ".csv",
    ".tsv",
}


def normalize_header(value: Any) -> str:
    text = "" if value is None else str(value)
    return unicodedata.normalize("NFKC", text).strip().casefold()


def json_cell(value: Any) -> Any:
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    return str(value)


@dataclass(frozen=True)
class ColumnDef:
    id: str
    name: str
    display_name: str
    normalized_name: str
    index: int
    occurrence: int
    key: str

    def as_dict(self) -> dict[str, Any]:
        return {
            "id": self.id,
            "name": self.name,
            "displayName": self.display_name,
            "normalizedName": self.normalized_name,
            "index": self.index,
            "occurrence": self.occurrence,
            "key": self.key,
        }


@dataclass
class SheetData:
    id: str
    name: str
    columns: list[ColumnDef]
    rows: list[list[Any]]
    declared_range: str | None
    actual_range: str

    @property
    def headers(self) -> list[str]:
        return [column.display_name for column in self.columns]

    def as_metadata(self) -> dict[str, Any]:
        return {
            "id": self.id,
            "name": self.name,
            "rowCount": len(self.rows),
            "columnCount": len(self.columns),
            "declaredRange": self.declared_range,
            "actualRange": self.actual_range,
            "rangeCorrected": bool(self.declared_range and self.declared_range != self.actual_range),
            "columns": [column.as_dict() for column in self.columns],
        }


@dataclass
class WorkbookData:
    sheets: list[SheetData]

    def sheet_by_id(self, sheet_id: str) -> SheetData:
        for sheet in self.sheets:
            if sheet.id == sheet_id:
                return sheet
        raise KeyError(sheet_id)

    def metadata(self) -> dict[str, Any]:
        coverage: dict[str, dict[str, Any]] = {}
        for sheet in self.sheets:
            seen: set[str] = set()
            for column in sheet.columns:
                if column.key in seen:
                    continue
                seen.add(column.key)
                item = coverage.setdefault(
                    column.key,
                    {
                        "key": column.key,
                        "headerName": column.name,
                        "displayName": column.display_name,
                        "normalizedName": column.normalized_name,
                        "occurrence": column.occurrence,
                        "fallbackColumnIndex": column.index,
                        "sheetIds": [],
                    },
                )
                item["sheetIds"].append(sheet.id)
        unified = []
        total_sheets = len(self.sheets)
        for item in coverage.values():
            item["coverage"] = len(item["sheetIds"])
            item["totalSheets"] = total_sheets
            unified.append(item)
        return {
            "sheetCount": total_sheets,
            "totalRows": sum(len(sheet.rows) for sheet in self.sheets),
            "maxColumns": max((len(sheet.columns) for sheet in self.sheets), default=0),
            "sheets": [sheet.as_metadata() for sheet in self.sheets],
            "unifiedColumns": unified,
        }


def _last_nonempty_index(row: Iterable[Any]) -> int:
    last = 0
    for index, value in enumerate(row, start=1):
        if value is not None and value != "":
            last = index
    return last


def _build_sheet(
    sheet_id: str,
    name: str,
    raw_rows: list[list[Any]],
    declared_range: str | None = None,
) -> SheetData:
    while raw_rows and _last_nonempty_index(raw_rows[-1]) == 0:
        raw_rows.pop()
    width = max((_last_nonempty_index(row) for row in raw_rows), default=0)
    width = max(width, 1)
    if not raw_rows:
        raw_rows = [[None] * width]
    padded = [list(row[:width]) + [None] * max(0, width - len(row)) for row in raw_rows]
    raw_header = padded[0]
    data_rows = [[json_cell(value) for value in row] for row in padded[1:]]

    names: list[str] = []
    normalized: list[str] = []
    for index, value in enumerate(raw_header, start=1):
        name_value = "" if value is None else str(value).strip()
        if not name_value:
            name_value = f"列{get_column_letter(index)}"
        names.append(name_value)
        normalized.append(normalize_header(name_value))

    totals: dict[str, int] = {}
    for value in normalized:
        totals[value] = totals.get(value, 0) + 1
    occurrences: dict[str, int] = {}
    columns: list[ColumnDef] = []
    for index, (name_value, normalized_value) in enumerate(zip(names, normalized), start=1):
        occurrence = occurrences.get(normalized_value, 0) + 1
        occurrences[normalized_value] = occurrence
        duplicate = totals[normalized_value] > 1
        display_name = f"{name_value}（{get_column_letter(index)}列）" if duplicate else name_value
        key = f"{normalized_value}#{occurrence}"
        columns.append(
            ColumnDef(
                id=f"{sheet_id}:c{index}",
                name=name_value,
                display_name=display_name,
                normalized_name=normalized_value,
                index=index,
                occurrence=occurrence,
                key=key,
            )
        )

    last_cell = f"{get_column_letter(width)}{len(padded)}"
    return SheetData(
        id=sheet_id,
        name=name,
        columns=columns,
        rows=data_rows,
        declared_range=declared_range,
        actual_range=f"A1:{last_cell}",
    )


class WorkbookReader:
    def read(self, path: Path) -> WorkbookData:
        extension = path.suffix.casefold()
        if extension not in SUPPORTED_EXTENSIONS:
            raise ValueError(f"不支持的文件格式：{extension or '无扩展名'}")
        if extension in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
            return self._read_openxml(path)
        if extension in {".xls", ".xlt"}:
            return self._read_xls(path)
        if extension == ".xlsb":
            return self._read_xlsb(path)
        if extension in {".csv", ".tsv"}:
            return self._read_delimited(path, "\t" if extension == ".tsv" else None)
        return self._read_spreadsheet_xml(path)

    def _read_openxml(self, path: Path) -> WorkbookData:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        sheets: list[SheetData] = []
        try:
            for index, worksheet in enumerate(workbook.worksheets, start=1):
                declared = worksheet.calculate_dimension(force=True)
                # Some WPS/Excel exports declare A1 although the XML contains thousands of rows.
                # Resetting dimensions makes the read-only parser scan the actual sheet data.
                worksheet.reset_dimensions()
                raw_rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
                sheets.append(_build_sheet(f"sheet-{index}", worksheet.title, raw_rows, declared))
        finally:
            workbook.close()
        return WorkbookData(sheets)

    def _read_xls(self, path: Path) -> WorkbookData:
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError("读取 .xls/.xlt 需要安装 xlrd") from exc

        workbook = xlrd.open_workbook(path, on_demand=True)
        sheets: list[SheetData] = []
        try:
            for index, worksheet in enumerate(workbook.sheets(), start=1):
                rows: list[list[Any]] = []
                for row_index in range(worksheet.nrows):
                    values: list[Any] = []
                    for column_index in range(worksheet.ncols):
                        cell = worksheet.cell(row_index, column_index)
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, workbook.datemode)
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            value = bool(value)
                        elif cell.ctype == xlrd.XL_CELL_EMPTY:
                            value = None
                        values.append(value)
                    rows.append(values)
                declared = f"A1:{get_column_letter(max(worksheet.ncols, 1))}{max(worksheet.nrows, 1)}"
                sheets.append(_build_sheet(f"sheet-{index}", worksheet.name, rows, declared))
        finally:
            workbook.release_resources()
        return WorkbookData(sheets)

    def _read_xlsb(self, path: Path) -> WorkbookData:
        try:
            from pyxlsb import open_workbook
        except ImportError as exc:
            raise RuntimeError("读取 .xlsb 需要安装 pyxlsb") from exc

        sheets: list[SheetData] = []
        with open_workbook(path) as workbook:
            for index, sheet_name in enumerate(workbook.sheets, start=1):
                with workbook.get_sheet(sheet_name) as worksheet:
                    rows = [[cell.v for cell in row] for row in worksheet.rows()]
                sheets.append(_build_sheet(f"sheet-{index}", sheet_name, rows))
        return WorkbookData(sheets)

    def _read_delimited(self, path: Path, delimiter: str | None) -> WorkbookData:
        text: str | None = None
        used_encoding = "utf-8-sig"
        for encoding in ("utf-8-sig", "gb18030", "utf-16"):
            try:
                text = path.read_text(encoding=encoding)
                used_encoding = encoding
                break
            except UnicodeError:
                continue
        if text is None:
            raise ValueError("无法识别 CSV/TSV 文件编码")
        if delimiter is None:
            try:
                delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
            except csv.Error:
                delimiter = ","
        rows = [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
        sheet_name = path.stem[:31] or "数据"
        sheet = _build_sheet("sheet-1", sheet_name, rows)
        sheet.declared_range = f"delimited:{used_encoding}:{repr(delimiter)}"
        return WorkbookData([sheet])

    def _read_spreadsheet_xml(self, path: Path) -> WorkbookData:
        namespace = "urn:schemas-microsoft-com:office:spreadsheet"
        ns = {"ss": namespace}
        tree = ET.parse(path)
        root = tree.getroot()
        sheets: list[SheetData] = []
        for sheet_index, worksheet in enumerate(root.findall(".//ss:Worksheet", ns), start=1):
            name = worksheet.attrib.get(f"{{{namespace}}}Name", f"Sheet{sheet_index}")
            table = worksheet.find("ss:Table", ns)
            raw_rows: list[list[Any]] = []
            if table is not None:
                for row in table.findall("ss:Row", ns):
                    values: list[Any] = []
                    for cell in row.findall("ss:Cell", ns):
                        target_index = int(cell.attrib.get(f"{{{namespace}}}Index", len(values) + 1))
                        while len(values) < target_index - 1:
                            values.append(None)
                        data = cell.find("ss:Data", ns)
                        value: Any = None
                        if data is not None:
                            value = data.text or ""
                            data_type = data.attrib.get(f"{{{namespace}}}Type", "String")
                            if data_type == "Number":
                                numeric = float(value)
                                value = int(numeric) if numeric.is_integer() else numeric
                            elif data_type == "Boolean":
                                value = value == "1"
                            elif data_type == "DateTime":
                                value = datetime.fromisoformat(value.rstrip("Z"))
                        values.append(value)
                    raw_rows.append(values)
            sheets.append(_build_sheet(f"sheet-{sheet_index}", name, raw_rows))
        if not sheets:
            raise ValueError("未在 XML 文件中找到 Excel Worksheet")
        return WorkbookData(sheets)
