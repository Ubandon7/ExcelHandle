from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="147D64")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _styled_header(sheet: Any, headers: list[str]) -> None:
    cells = []
    for value in headers:
        cell = WriteOnlyCell(sheet, value=value)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        cells.append(cell)
    sheet.append(cells)
    sheet.freeze_panes = "A2"
    if headers:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
        for index, header in enumerate(headers, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = min(max(len(str(header)) + 3, 12), 32)


def _safe_row(sheet: Any, values: Iterable[Any]) -> list[Any]:
    output: list[Any] = []
    for value in values:
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            cell = WriteOnlyCell(sheet, value=value)
            cell.data_type = "s"
            output.append(cell)
        else:
            output.append(value)
    return output


def export_filter_workbook(result: dict[str, Any], destination: Path) -> None:
    workbook = Workbook(write_only=True)
    for key in result["sheetOrder"]:
        dataset = result["sheets"][key]
        sheet = workbook.create_sheet(dataset["name"])
        _styled_header(sheet, dataset["headers"])
        for row in dataset["rows"]:
            sheet.append(_safe_row(sheet, row))
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)


def export_sum_workbook(result: dict[str, Any], destination: Path) -> None:
    workbook = Workbook(write_only=True)
    summary = workbook.create_sheet("汇总结果")
    summary_headers = [
        "Sheet 名",
        "目标列",
        "数据行数",
        "参与求和数值数",
        "文本转换数",
        "空值数",
        "非数值数",
        "错误值数",
        "求和结果",
        "状态",
        "说明",
    ]
    _styled_header(summary, summary_headers)
    for item in result["results"]:
        sum_value = float(item["sum"]) if item["sum"] is not None else None
        summary.append(
            [
                item["sheetName"],
                item["columnName"],
                item["dataRowCount"],
                item["numericCount"],
                item["textConvertedCount"],
                item["emptyCount"],
                item["nonNumericCount"],
                item["errorValueCount"],
                sum_value,
                item["status"],
                item["message"],
            ]
        )
    summary.append(["全部 Sheet 总计", "", "", "", "", "", "", "", float(result["grandTotal"]), "", ""])

    details = workbook.create_sheet("异常明细")
    detail_headers = ["来源 Sheet", "源行号", "列序号", "目标列", "原始值", "未参与原因"]
    _styled_header(details, detail_headers)
    for item in result["invalidDetails"]:
        details.append(
            _safe_row(
                details,
                [
                    item["sheetName"],
                    item["sourceRow"],
                    item["columnIndex"],
                    item["columnName"],
                    item["originalValue"],
                    item["reason"],
                ],
            )
        )

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)

