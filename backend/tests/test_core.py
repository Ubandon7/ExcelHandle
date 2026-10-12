from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from app.exports import export_filter_workbook, export_sum_workbook
from app.filters import execute_affix_split, execute_column_split, execute_filter_groups
from app.sums import execute_sum, parse_number
from app.workbook import WorkbookData, WorkbookReader, _build_sheet


def make_workbook() -> WorkbookData:
    return WorkbookData(
        sheets=[
            _build_sheet(
                "sheet-1",
                "华东",
                [
                    ["客户", "区域", "金额"],
                    ["Alpha", "上海", "$1,200.50"],
                    ["Beta", "杭州", "12%"],
                    ["Gamma", "上海", "约12元"],
                    ["Delta", "南京", -100],
                ],
                "A1:C5",
            ),
            _build_sheet(
                "sheet-2",
                "华南",
                [
                    ["客户", "区域", "金额", "备注"],
                    ["Alpine", "深圳", "(100)", "重点"],
                    ["Echo", "广州", "", ""],
                ],
                "A1:D3",
            ),
        ]
    )


def condition(header: str, operator: str, value: str | None = None) -> dict:
    return {
        "type": "condition",
        "columnRef": {"headerName": header, "fallbackColumnIndex": 1, "occurrence": 1},
        "operator": operator,
        "value": value,
        "values": [],
        "options": {"caseSensitive": False, "trimCell": False, "trimInput": True},
    }


def test_numeric_parsing_business_rules() -> None:
    assert parse_number("$1,234.50") == (Decimal("1234.50"), "text_converted")
    assert parse_number("12%") == (Decimal("12"), "text_converted")
    assert parse_number("(100)") == (Decimal("100"), "text_converted")
    assert parse_number("-100") == (Decimal("-100"), "text_converted")
    assert parse_number("约12元") == (None, "contains_text")
    assert parse_number("3件") == (None, "contains_text")


def test_filter_nested_or_first_match_union_and_unmatched() -> None:
    workbook = make_workbook()
    groups = [
        {
            "id": "g1",
            "name": "A开头或上海",
            "priority": 1,
            "sourceScope": {"mode": "ALL", "sheetId": None},
            "conditionTree": {
                "type": "group",
                "relation": "OR",
                "children": [
                    condition("客户", "starts_with", "Al"),
                    condition("区域", "in") | {"values": ["上海"]},
                ],
            },
        },
        {
            "id": "g2",
            "name": "华南金额非空",
            "priority": 2,
            "sourceScope": {"mode": "SINGLE", "sheetId": "sheet-2"},
            "conditionTree": {
                "type": "group",
                "relation": "AND",
                "children": [condition("金额", "not_empty")],
            },
        },
    ]

    result = execute_filter_groups(workbook, groups)

    assert result["sourceRowCount"] == 6
    assert result["assignedRowCount"] == 6
    assert result["rowCountConserved"] is True
    assert result["sheetOrder"] == ["g1", "g2", "__unmatched__"]
    assert result["sheets"]["g1"]["rowCount"] == 3
    assert result["sheets"]["g2"]["rowCount"] == 0  # Alpine first matched g1 and cannot repeat.
    assert result["sheets"]["__unmatched__"]["rowCount"] == 3
    assert result["sheets"]["g1"]["headers"] == ["来源 Sheet", "客户", "区域", "金额", "备注"]
    assert result["sheets"]["__unmatched__"]["isSystem"] is True


def test_column_split_groups_unique_values_and_sanitizes_sheet_names() -> None:
    workbook = WorkbookData(sheets=[
        _build_sheet("sheet-1", "数据", [
            ["分类", "编号"],
            ["A/B", 1],
            ["A:B", 2],
            [None, 3],
            [" ", 4],
            ["空白", 5],
            ["x" * 40, 6],
            ["x" * 39 + "y", 7],
        ]),
        _build_sheet("sheet-2", "无分类列", [["编号"], [8]]),
    ])

    result = execute_column_split(
        workbook,
        {"headerName": "分类", "occurrence": 1, "fallbackColumnIndex": 1},
        {"mode": "ALL", "sheetId": None},
    )

    assert result["sourceRowCount"] == result["assignedRowCount"] == 8
    assert result["rowCountConserved"] is True
    assert [result["sheets"][key]["name"] for key in result["sheetOrder"]] == [
        "A_B", "A_B_2", "空白", "空白_2", "x" * 31, f"{'x' * 29}_2", "未匹配",
    ]
    assert result["sheets"]["category-3"]["rowCount"] == 2
    assert result["sheets"]["__unmatched__"]["rowCount"] == 1


def test_affix_split_uses_first_prefix_and_last_suffix_delimiter() -> None:
    workbook = WorkbookData(sheets=[_build_sheet("sheet-1", "编码", [
        ["名称"],
        ["A-001"],
        ["A-002"],
        ["B-001"],
        ["A-middle-003"],
        ["plain"],
        ["-root"],
        ["tail-"],
    ])])
    column_ref = {"headerName": "名称", "occurrence": 1, "fallbackColumnIndex": 1}
    scope = {"mode": "ALL", "sheetId": None}

    prefixes = execute_affix_split(workbook, column_ref, scope, "PREFIX", "-")
    assert [(prefixes["sheets"][key]["name"], prefixes["sheets"][key]["rowCount"]) for key in prefixes["sheetOrder"]] == [
        ("A-", 3), ("B-", 1), ("tail-", 1), ("未匹配", 2),
    ]
    suffixes = execute_affix_split(workbook, column_ref, scope, "SUFFIX", "-")
    assert [(suffixes["sheets"][key]["name"], suffixes["sheets"][key]["rowCount"]) for key in suffixes["sheetOrder"]] == [
        ("-001", 2), ("-002", 1), ("-003", 1), ("-root", 1), ("未匹配", 2),
    ]
    assert prefixes["rowCountConserved"] is suffixes["rowCountConserved"] is True


def test_no_prefix_and_no_suffix_include_missing_or_empty_affixes() -> None:
    workbook = WorkbookData(sheets=[_build_sheet(
        "sheet-affix",
        "标识符规则",
        [
            ["名称"],
            ["Alpha"],
            ["-root"],
            ["name-"],
            ["xj-name"],
            ["xj-name-admin"],
        ],
        "A1:A6",
    )])
    groups = [
        {
            "id": "no-prefix", "name": "无前缀", "priority": 1,
            "sourceScope": {"mode": "ALL", "sheetId": None},
            "conditionTree": {"type": "group", "relation": "AND", "children": [condition("名称", "no_prefix", "-")]},
        },
        {
            "id": "no-suffix", "name": "无后缀", "priority": 2,
            "sourceScope": {"mode": "ALL", "sheetId": None},
            "conditionTree": {"type": "group", "relation": "AND", "children": [condition("名称", "no_suffix", "-")]},
        },
    ]

    result = execute_filter_groups(workbook, groups)

    assert result["sheets"]["no-prefix"]["rows"] == [["标识符规则", "Alpha"], ["标识符规则", "-root"]]
    assert result["sheets"]["no-suffix"]["rows"] == [["标识符规则", "name-"]]
    assert result["sheets"]["__unmatched__"]["rows"] == [
        ["标识符规则", "xj-name"],
        ["标识符规则", "xj-name-admin"],
    ]


def test_sum_by_header_and_fixed_position() -> None:
    workbook = make_workbook()
    by_header = execute_sum(workbook, {"matchMode": "NORMALIZED_HEADER", "headerName": " 金额 "})
    assert by_header["grandTotal"] == "1212.50"
    assert by_header["invalidCount"] == 1
    assert by_header["invalidDetails"][0]["originalValue"] == "约12元"
    assert by_header["results"][0]["textConvertedCount"] == 2
    assert by_header["results"][0]["numericCount"] == 3

    by_position = execute_sum(workbook, {"matchMode": "FIXED_POSITION", "columnIndex": 3})
    assert by_position["grandTotal"] == "1212.50"


def test_exports_are_valid_xlsx(tmp_path: Path) -> None:
    workbook = make_workbook()
    groups = [{
        "id": "g1", "name": "上海", "priority": 1,
        "sourceScope": {"mode": "ALL", "sheetId": None},
        "conditionTree": {"type": "group", "relation": "AND", "children": [condition("区域", "in") | {"values": ["上海"]}]},
    }]
    filter_result = execute_filter_groups(workbook, groups)
    filter_path = tmp_path / "filter.xlsx"
    export_filter_workbook(filter_result, filter_path)
    filter_book = load_workbook(filter_path, read_only=True, data_only=True)
    assert filter_book.sheetnames == ["上海", "未匹配"]
    assert sum(1 for _ in filter_book["上海"].iter_rows()) == 3
    assert sum(1 for _ in filter_book["未匹配"].iter_rows()) == 5

    sum_result = execute_sum(workbook, {"matchMode": "NORMALIZED_HEADER", "headerName": "金额"})
    sum_path = tmp_path / "sum.xlsx"
    export_sum_workbook(sum_result, sum_path)
    sum_book = load_workbook(sum_path, read_only=True, data_only=True)
    assert sum_book.sheetnames == ["汇总结果", "异常明细"]
    assert sum_book["汇总结果"]["I4"].value == 1212.5
    assert sum_book["异常明细"]["E2"].value == "约12元"


def test_reference_samples_have_expected_conservation() -> None:
    project_root = Path(__file__).resolve().parents[2]
    source = WorkbookReader().read(project_root / "样例" / "原始全部数据.xlsx")
    expected = WorkbookReader().read(project_root / "样例" / "根据不同表头及条件筛选后的结果.xlsx")
    source_rows = sum(len(sheet.rows) for sheet in source.sheets)
    expected_rows = sum(len(sheet.rows) for sheet in expected.sheets)
    assert len(expected.sheets) == 73
    assert source_rows == expected_rows == 6991


def test_csv_preserves_quoted_line_breaks(tmp_path: Path) -> None:
    path = tmp_path / "带换行.csv"
    path.write_text('名称,备注\nA001,"第一行\n第二行"\n', encoding="utf-8")
    workbook = WorkbookReader().read(path)
    assert workbook.sheets[0].rows == [["A001", "第一行\n第二行"]]
