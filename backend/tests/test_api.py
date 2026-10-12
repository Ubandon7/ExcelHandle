from __future__ import annotations

from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.main import app
from app.services import ApplicationService
from app.storage import Repository


def workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "一月"
    sheet.append(["名称", "分类", "金额"])
    sheet.append(["A-001", "甲", "$10"])
    sheet.append(["B-001", "乙", "约12元"])
    second = workbook.create_sheet("二月")
    second.append(["名称", "分类", "金额"])
    second.append(["A-002", "甲", "(5)"])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_complete_http_workflow(tmp_path: Path) -> None:
    app.state.service = ApplicationService(Repository(tmp_path / "data"))
    with TestClient(app) as client:
        upload = client.post("/api/files", files={"file": ("demo.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert upload.status_code == 201
        file = upload.json()["file"]
        assert file["metadata"]["sheetCount"] == 2
        assert file["metadata"]["totalRows"] == 3

        files = client.get("/api/files")
        assert files.status_code == 200
        assert [item["id"] for item in files.json()["files"]] == [file["id"]]

        prefixes = client.get(f"/api/files/{file['id']}/affix-suggestions", params={
            "headerName": "名称", "occurrence": 1, "kind": "PREFIX", "delimiter": "-", "scopeMode": "ALL",
        })
        assert prefixes.status_code == 200
        assert prefixes.json()["items"] == [{"value": "A-", "count": 2}, {"value": "B-", "count": 1}]
        suffixes = client.get(f"/api/files/{file['id']}/affix-suggestions", params={
            "headerName": "名称", "occurrence": 1, "kind": "SUFFIX", "delimiter": "-", "scopeMode": "ALL",
        })
        assert suffixes.status_code == 200
        assert suffixes.json()["items"] == [{"value": "-001", "count": 2}, {"value": "-002", "count": 1}]

        group_payload = {
            "name": "A类",
            "sourceScope": {"mode": "ALL", "sheetId": None},
            "conditionTree": {
                "type": "group", "relation": "AND",
                "children": [{
                    "type": "condition",
                    "columnRef": {"headerName": "名称", "fallbackColumnIndex": 1, "occurrence": 1},
                    "operator": "starts_with", "value": "A", "values": [],
                    "options": {"caseSensitive": False, "trimCell": False, "trimInput": True},
                }],
            },
        }
        created = client.post(f"/api/files/{file['id']}/filter-groups", json=group_payload)
        assert created.status_code == 201
        group = created.json()["group"]
        assert group["saved"] is False

        saved_group = client.patch(f"/api/filter-groups/{group['id']}", json={"saved": True})
        assert saved_group.status_code == 200
        group = saved_group.json()["group"]
        assert group["saved"] is True
        assert group["savedAt"]

        duplicate = client.post(f"/api/files/{file['id']}/filter-groups", json=group_payload)
        assert duplicate.status_code == 400
        assert "重复" in duplicate.json()["detail"]

        filter_response = client.post(f"/api/files/{file['id']}/filter-runs", json={"orderedGroupIds": [group["id"]]})
        assert filter_response.status_code == 202
        filter_run = client.get(f"/api/filter-runs/{filter_response.json()['run']['id']}").json()["run"]
        assert filter_run["status"] == "succeeded"
        assert filter_run["saved"] is False
        assert filter_run["sourceRowCount"] == filter_run["assignedRowCount"] == 3
        assert [item["rowCount"] for item in filter_run["summary"]] == [2, 1]

        sum_sources = client.get("/api/sum-sources")
        assert sum_sources.status_code == 200
        sources = sum_sources.json()["sources"]
        assert [(item["type"], item["id"]) for item in sources] == [
            ("FILTER_RESULT", filter_run["id"]),
            ("FILE", file["id"]),
        ]
        assert sources[0]["saved"] is False
        assert sources[0]["metadata"]["sheetCount"] == 2

        filtered_sum_response = client.post("/api/sum-runs", json={
            "sourceRef": {"type": "FILTER_RESULT", "id": filter_run["id"]},
            "columnSelector": {"matchMode": "NORMALIZED_HEADER", "headerName": "金额"},
            "numberPolicy": {},
        })
        assert filtered_sum_response.status_code == 202
        filtered_sum = client.get(f"/api/sum-runs/{filtered_sum_response.json()['run']['id']}").json()["run"]
        assert filtered_sum["status"] == "succeeded"
        assert filtered_sum["grandTotal"] == "15"
        assert filtered_sum["invalidCount"] == 1

        filter_export = client.post(f"/api/filter-runs/{filter_run['id']}/exports")
        download = client.get(filter_export.json()["downloadUrl"])
        assert download.status_code == 200
        exported = load_workbook(BytesIO(download.content), read_only=True)
        assert exported.sheetnames == ["A类", "未匹配"]
        repeated_export = client.post(f"/api/filter-runs/{filter_run['id']}/exports")
        assert repeated_export.status_code == 200
        assert repeated_export.json() == filter_export.json()

        history = client.get(f"/api/files/{file['id']}/filter-runs")
        assert history.status_code == 200
        assert history.json()["runs"] == []
        temporary = client.get(f"/api/files/{file['id']}/temporary-filter-run")
        assert temporary.status_code == 200
        assert temporary.json()["run"]["id"] == filter_run["id"]

        saved = client.post(f"/api/filter-runs/{filter_run['id']}/save")
        assert saved.status_code == 200
        assert saved.json()["run"]["saved"] is True
        saved_sources = client.get("/api/sum-sources").json()["sources"]
        assert next(item for item in saved_sources if item["id"] == filter_run["id"])["saved"] is True
        assert client.get(f"/api/files/{file['id']}/temporary-filter-run").json()["run"] is None
        history = client.get(f"/api/files/{file['id']}/filter-runs")
        assert [item["id"] for item in history.json()["runs"]] == [filter_run["id"]]

        removed = client.delete(f"/api/filter-runs/{filter_run['id']}")
        assert removed.status_code == 204
        assert client.get(f"/api/filter-runs/{filter_run['id']}").status_code == 404
        assert client.get(f"/api/files/{file['id']}/filter-runs").json()["runs"] == []

        sum_response = client.post(f"/api/files/{file['id']}/sum-runs", json={
            "columnSelector": {"matchMode": "NORMALIZED_HEADER", "headerName": "金额"},
            "numberPolicy": {},
        })
        assert sum_response.status_code == 202
        sum_run = client.get(f"/api/sum-runs/{sum_response.json()['run']['id']}").json()["run"]
        assert sum_run["status"] == "succeeded"
        assert sum_run["grandTotal"] == "15"
        assert sum_run["invalidCount"] == 1
        invalid = client.get(f"/api/sum-runs/{sum_run['id']}/invalid-details").json()["items"]
        assert invalid[0]["originalValue"] == "约12元"


def test_invalid_workbook_returns_readable_400(tmp_path: Path) -> None:
    app.state.service = ApplicationService(Repository(tmp_path / "data"))
    with TestClient(app) as client:
        response = client.post("/api/files", files={"file": ("broken.xlsx", b"not an excel file")})
        assert response.status_code == 400
        assert "无法解析" in response.json()["detail"]


def test_column_split_http_workflow(tmp_path: Path) -> None:
    app.state.service = ApplicationService(Repository(tmp_path / "data"))
    with TestClient(app) as client:
        file = client.post(
            "/api/files",
            files={"file": ("demo.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        ).json()["file"]
        response = client.post(f"/api/files/{file['id']}/column-split-runs", json={
            "columnRef": {"headerName": "分类", "occurrence": 1, "fallbackColumnIndex": 2},
            "sourceScope": {"mode": "ALL", "sheetId": None},
        })

        assert response.status_code == 202
        run = client.get(f"/api/filter-runs/{response.json()['run']['id']}").json()["run"]
        assert run["status"] == "succeeded"
        assert run["sourceRowCount"] == run["assignedRowCount"] == 3
        assert [(item["name"], item["rowCount"]) for item in run["summary"]] == [("甲", 2), ("乙", 1)]

        exported_response = client.post(f"/api/filter-runs/{run['id']}/exports")
        exported = client.get(exported_response.json()["downloadUrl"])
        workbook = load_workbook(BytesIO(exported.content), read_only=True)
        assert workbook.sheetnames == ["甲", "乙"]

        affix_response = client.post(f"/api/files/{file['id']}/affix-split-runs", json={
            "columnRef": {"headerName": "名称", "occurrence": 1, "fallbackColumnIndex": 1},
            "sourceScope": {"mode": "ALL", "sheetId": None},
            "kind": "PREFIX",
            "delimiter": "-",
        })
        assert affix_response.status_code == 202
        affix_run = client.get(f"/api/filter-runs/{affix_response.json()['run']['id']}").json()["run"]
        assert affix_run["status"] == "succeeded"
        assert [(item["name"], item["rowCount"]) for item in affix_run["summary"]] == [("A-", 2), ("B-", 1)]


def test_saved_conditions_survive_clear_and_workbook_deletion(tmp_path: Path) -> None:
    repository = Repository(tmp_path / "data")
    app.state.service = ApplicationService(repository)
    with TestClient(app) as client:
        uploaded = client.post(
            "/api/files",
            files={"file": ("persistent.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        ).json()["file"]
        payload = {
            "name": "永久条件",
            "sourceScope": {"mode": "ALL", "sheetId": None},
            "conditionTree": {
                "type": "group", "relation": "AND",
                "children": [{
                    "type": "condition",
                    "columnRef": {"headerName": "名称", "fallbackColumnIndex": 1, "occurrence": 1},
                    "operator": "starts_with", "value": "A", "values": [],
                    "options": {"caseSensitive": False, "trimCell": False, "trimInput": True},
                }],
            },
        }
        saved_group = client.post(f"/api/files/{uploaded['id']}/filter-groups", json=payload).json()["group"]
        saved_group = client.patch(f"/api/filter-groups/{saved_group['id']}", json={"saved": True}).json()["group"]
        unsaved_group = client.post(
            f"/api/files/{uploaded['id']}/filter-groups",
            json=payload | {"name": "临时条件"},
        ).json()["group"]

        global_groups = client.get("/api/filter-groups/saved")
        assert global_groups.status_code == 200
        assert [item["id"] for item in global_groups.json()["groups"]] == [saved_group["id"]]

        other_file = client.post(
            "/api/files",
            files={"file": ("other.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        ).json()["file"]
        cloned = client.post(
            f"/api/files/{other_file['id']}/filter-groups/from-saved/{saved_group['id']}"
        )
        assert cloned.status_code == 201
        cloned_group = cloned.json()["group"]
        assert cloned_group["fileId"] == other_file["id"]
        assert cloned_group["saved"] is False
        assert cloned_group["conditionTree"] == saved_group["conditionTree"]

        cleared = client.delete(f"/api/files/{uploaded['id']}/filter-groups/current")
        assert cleared.status_code == 200
        assert cleared.json()["clearedCount"] == 2
        groups = client.get(f"/api/files/{uploaded['id']}/filter-groups").json()["groups"]
        assert [(item["name"], item["saved"], item["active"]) for item in groups] == [("永久条件", True, False)]
        assert repository.get_group(unsaved_group["id"]) is None

        assert client.delete(f"/api/files/{uploaded['id']}").status_code == 204
        orphan = repository.get_group(saved_group["id"])
        assert orphan is not None
        assert orphan["saved"] is True
        assert orphan["active"] is False
        assert orphan["fileId"] is None

        replacement = client.post(
            "/api/files",
            files={"file": ("persistent.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        ).json()["file"]
        restored = client.get(f"/api/files/{replacement['id']}/filter-groups").json()["groups"]
        assert [(item["id"], item["active"]) for item in restored] == [(saved_group["id"], False)]

        assert client.delete(f"/api/filter-groups/{saved_group['id']}").status_code == 204
        assert repository.get_group(saved_group["id"]) is None
